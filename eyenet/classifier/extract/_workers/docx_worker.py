"""In-jail DOCX text extractor (python-docx). Runs ONLY inside nsjail.

Imports ``docx`` from the eyenet-extract venv (bound RO, on PYTHONPATH). Reads
/input, gathers paragraph text plus table-cell text, and emits the
{"text", "meta"} envelope. A non-DOCX zip (or corrupt file) raises out of
docx.Document → non-zero exit → fail closed.

M10 slice 7: also surfaces the OPC core properties (``core.xml``) under
``meta.embedded`` — creator/lastModifiedBy/created/modified/revision/… — often
the most probative attribution fields. ATTACKER-CONTROLLED: captured verbatim,
sanitized HOST-SIDE. ``created``/``modified``/``last_printed`` are ``datetime``
objects, so they are rendered to ISO strings here — emitting a raw ``datetime``
would make ``json.dumps`` raise and fail-close an otherwise-readable document.
"""

import io
import json
import sys

import docx


def main() -> int:
    with open("/input", "rb") as handle:
        data = handle.read()
    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    text = "\n".join(part for part in parts if part)

    core = document.core_properties
    embedded = {
        "author": core.author,
        "category": core.category,
        "comments": core.comments,
        "content_status": core.content_status,
        "created": core.created.isoformat() if core.created else None,
        "identifier": core.identifier,
        "keywords": core.keywords,
        "language": core.language,
        "last_modified_by": core.last_modified_by,
        "last_printed": core.last_printed.isoformat() if core.last_printed else None,
        "modified": core.modified.isoformat() if core.modified else None,
        "revision": core.revision,
        "subject": core.subject,
        "title": core.title,
        "version": core.version,
    }
    embedded = {key: value for key, value in embedded.items() if value}

    meta = {"method": "python-docx", "embedded": embedded}
    sys.stdout.write(json.dumps({"text": text, "meta": meta}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
